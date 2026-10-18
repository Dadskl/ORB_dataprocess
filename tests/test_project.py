"""Portable smoke tests for archive discovery and optional Excel export."""
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from input_sources import collect_inputs
from origin_export import write_workbook


class InputDiscoveryTests(unittest.TestCase):
    def test_archives_with_identical_inner_paths_stay_independent(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name in ('mono.zip', 'stereo.zip'):
                with ZipFile(root / name, 'w') as archive:
                    archive.writestr('results/sequence_00/statistics/frames.csv', 'timestamp\n0\n')
            runs, failures = collect_inputs(root, root / 'extracted')
            self.assertFalse(failures)
            self.assertEqual(len(runs), 2)
            self.assertNotEqual(runs[0][0], runs[1][0])
            self.assertTrue(all(run.joinpath('statistics/frames.csv').is_file() for run, _ in runs))

    def test_unsafe_zip_is_reported_without_hiding_valid_sibling(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with ZipFile(root / 'bad.zip', 'w') as archive:
                archive.writestr('../escape.txt', 'bad')
            with ZipFile(root / 'good.zip', 'w') as archive:
                archive.writestr('sequence_00/statistics/frames.csv', 'timestamp\n0\n')
            runs, failures = collect_inputs(root, root / 'extracted')
            self.assertEqual(len(runs), 1)
            self.assertEqual(len(failures), 1)
            self.assertFalse((root / 'escape.txt').exists())


class WorkbookTests(unittest.TestCase):
    def test_workbook_preserves_numeric_samples_and_sequence_ids(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest('optional openpyxl dependency is not installed')
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'result.xlsx'
            long_key = 'very_long_table_name_for_a_figure'
            tables = [
                {'key': 'PlotCatalog', 'headers': ['table_key', 'plot_type'],
                 'rows': [[long_key, 'Line']]},
                {'key': long_key, 'headers': ['sequence', 'error_m', 'source'],
                 'rows': [['00', 0.125, '=1+1'], ['01', None, 'recorder']]},
            ]
            write_workbook(tables, path)
            book = load_workbook(path, read_only=True, data_only=False)
            self.assertEqual(len(book.sheetnames), 2)
            sheet_name = book.sheetnames[1]
            self.assertLessEqual(len(sheet_name), 31)
            self.assertEqual(book['PlotCatalog']['A2'].value, sheet_name)
            self.assertEqual(book[sheet_name]['A2'].value, '00')
            self.assertEqual(book[sheet_name]['B2'].value, 0.125)
            self.assertEqual(book[sheet_name]['C2'].data_type, 's')
            self.assertIsNone(book[sheet_name]['B3'].value)
            book.close()


if __name__ == '__main__':
    unittest.main()
