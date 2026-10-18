"""Publication figures, CSV tables, HTML reports and optional Excel workbooks."""
from __future__ import annotations
import csv
import hashlib
import html
import json
from pathlib import Path
import re

import numpy as np
import matplotlib.pyplot as plt


def clean(value):
    if isinstance(value, np.generic): value = value.item()
    if isinstance(value, float) and not np.isfinite(value): return None
    if isinstance(value, Path): return str(value)
    return value


def runtime():
    """Check the optional, portable Excel dependency before processing inputs."""
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise ValueError('Excel export requires openpyxl. Install requirements-excel.txt.') from exc


def write_workbook(tables, path):
    """Stream all numeric tables to a standalone Excel workbook."""
    runtime()
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)
    names = {}
    for table in tables:
        key = table['key']
        safe = re.sub(r'[\\/\[\]:?*]', '_', key)
        name = safe if len(safe) <= 31 else safe[:24] + '_' + hashlib.sha256(key.encode()).hexdigest()[:6]
        if name in names.values():
            raise ValueError(f'Duplicate Excel worksheet name: {name}')
        names[key] = name
    header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='193B55')
    for table in tables:
        sheet = workbook.create_sheet(names[table['key']])
        sheet.freeze_panes = 'A2'
        sheet.sheet_properties.tabColor = '197278' if 'Summary' in table['key'] else '4679A4'
        headers = (['worksheet'] + table['headers']) if table['key'] == 'PlotCatalog' else table['headers']
        for i in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 24
        if table['key'] == 'ReadMe':
            sheet.column_dimensions['B'].width = 100
        header = []
        for value in headers:
            cell = WriteOnlyCell(sheet, value=value)
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(wrap_text=True)
            header.append(cell)
        sheet.append(header)
        for row in table['rows']:
            values = ([names.get(row[0], '')] + row) if table['key'] == 'PlotCatalog' else row
            cells = []
            for column, value in enumerate(values):
                # A source string must never become an Excel formula.
                if isinstance(value, str) and value.startswith('='):
                    value = "'" + value
                cell = WriteOnlyCell(sheet, value=value)
                if (table['key'].endswith('Summary') and column == 0) or isinstance(value, str):
                    cell.number_format = '@'
                elif isinstance(value, (int, float)):
                    cell.number_format = '0' if isinstance(value, int) else '0.000000'
                cells.append(cell)
            sheet.append(cells)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


class Report:
    def __init__(self, path, formats, dpi, excel=True):
        self.path, self.formats, self.dpi = Path(path), formats, max(1201, dpi)
        self.path.mkdir(parents=True, exist_ok=True)
        (self.path/'01_figures_hd').mkdir(exist_ok=True)
        (self.path/'02_origin').mkdir(exist_ok=True)
        self.images, self.metrics, self.notes, self.tables = [], {}, [], {}
        self.catalog = []
        self.excel = excel

    def save(self, name, fig):
        width, height = fig.get_size_inches()
        fig.set_size_inches(6.4, 6.4*height/width)
        for text in fig.findobj(match=lambda x: isinstance(x, plt.Text)):
            text.set_fontsize(min(6.5, text.get_fontsize()))
        fig.tight_layout(rect=getattr(fig, '_slam_layout_rect', (0, 0, 1, 1)))
        for ext in self.formats:
            kwargs = {'pil_kwargs': {'compression': 'tiff_lzw'}} if ext == 'tiff' else {}
            fig.savefig(self.path/'01_figures_hd'/f'{name}.{ext}', dpi=self.dpi, **kwargs)
        self.images.append(name)
        plt.close(fig)

    def save_chart(self, name, fig):
        """Rich diagnostic panels share the publication export settings."""
        self.save(name, fig)

    def add(self, name, headers, rows, plot='Data table', x='', y='', note=''):
        if len(headers) > 16384: raise ValueError('Too many Excel columns.')
        rows = [[clean(v) for v in row] for row in rows]
        if any(len(row) != len(headers) for row in rows): raise ValueError(f'Nonrectangular table: {name}')
        # Never silently truncate Excel's row limit; split large detail tables.
        chunks = max(1, (len(rows)+1048574)//1048575)
        for chunk in range(chunks):
            key = name if chunks == 1 else f'{name}_part{chunk+1}'
            part = rows[chunk*1048575:(chunk+1)*1048575]
            self.tables[key] = {'key': key, 'headers': list(headers), 'rows': part}
            self.catalog = [r for r in self.catalog if r[0] != key]
            self.catalog.append([key, plot, x, y, note])

    def lines(self, name, x, ys, ylabel, xlabel='Time since start (s)'):
        headers = [xlabel]+[f'{label} [{ylabel}]' for label in ys]
        self.add(name, headers, zip(x, *ys.values()), 'Line / scatter', headers[0], '; '.join(headers[1:]))

    def distribution(self, name, x, unit):
        x = np.asarray(x); x = x[np.isfinite(x)]
        if not len(x): return
        ordered = np.sort(x)
        self.add(name, [unit, 'ECDF_ratio'], zip(ordered, np.arange(1, len(x)+1)/len(x)),
                 'Box / histogram / ECDF', unit, 'ECDF_ratio', 'Box/histogram uses the first column; ECDF uses both columns.')
        counts, edges = np.histogram(x, bins=min(60, max(5, int(np.sqrt(len(x))))))
        self.add(name+'_hist', ['bin_left', 'bin_right', 'bin_center', 'count', 'fraction'],
                 zip(edges[:-1], edges[1:], (edges[:-1]+edges[1:])/2, counts, counts/len(x)),
                 'Histogram', 'bin_center', 'count', f'X unit: {unit}; final bin includes right edge.')

    def csv(self, name, headers, rows):
        rows = list(rows)
        with (self.path/'02_origin'/f'{name}.csv').open('w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f); writer.writerow(headers); writer.writerows(rows)
        self.add(name, headers, rows)

    def finish(self, title):
        (self.path/'metrics.json').write_text(json.dumps(self.metrics, indent=2, allow_nan=False, ensure_ascii=False)+'\n')
        flat = []
        def walk(prefix, value):
            if isinstance(value, dict):
                for k, v in value.items(): walk(f'{prefix}.{k}' if prefix else k, v)
            elif isinstance(value, (list, tuple)):
                for i, item in enumerate(value): walk(f'{prefix}[{i}]', item)
            else: flat.append([prefix, clean(value)])
        walk('', self.metrics)
        self.add('Metrics', ['metric', 'value'], flat)
        self.add('ReadMe', ['item', 'description'], [
            ['Source', str(self.metrics.get('source', title))],
            ['Data', 'Figures are rendered directly. CSV tables are in 03_data; optional Excel has the same numeric data.'],
            ['Missing', 'Empty numeric cells mean unavailable, never zero. Sequence IDs are text.'],
            ['Units', 'Units in headers; percent columns contain percentage points (0.81 means 0.81%).'],
            ['Method', 'Computed numeric snapshots; rerun Python after inputs change. No Excel formulas.'],
            ['Coordinates', 'Unaligned monocular position/speed uses arbitrary map units.'],
            ['Timing', '1000/tracking_ms is call capacity, NOT actual system FPS.'],
            ['Uncertainty', 'Per-frame STD is dispersion, not uncertainty across independent repeated experiments.'],
            *[['Note', n] for n in self.notes]])
        catalog = {'key': 'PlotCatalog', 'headers': ['table_key', 'plot_type', 'X', 'Y', 'note'], 'rows': self.catalog.copy()}
        tables = [catalog]+list(self.tables.values())
        if self.excel:
            write_workbook(tables, self.path/'02_origin/origin_data.xlsx')
        data_dir = self.path/'03_data'
        data_dir.mkdir(exist_ok=True)
        for table in tables:
            with (data_dir/(table['key']+'.csv')).open('w', newline='', encoding='utf-8-sig') as f:
                writer = csv.writer(f); writer.writerow(table['headers']); writer.writerows(table['rows'])
        (self.path/'figure_manifest.json').write_text(json.dumps({'figures': self.images, 'dpi': self.dpi, 'formats': self.formats}, indent=2)+'\n')
        cards = ''.join(f'<section><h2>{html.escape(n)}</h2><a href="01_figures_hd/{n}.png"><img loading="lazy" src="01_figures_hd/{n}.png"></a></section>' for n in self.images)
        notes = ''.join(f'<li>{html.escape(n)}</li>' for n in self.notes)
        (self.path/'report.html').write_text(f'''<!doctype html><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>body{{font:16px system-ui;margin:30px auto;max-width:1100px;color:#183047;background:#f3f5f8}}section{{background:white;padding:20px;margin:20px 0}}img{{width:100%}}</style>
<h1>{html.escape(title)}</h1><p>{'<a href="02_origin/origin_data.xlsx">Excel 数据</a> · ' if self.excel else ''}<a href="metrics.json">JSON 指标</a> · <a href="03_data/PlotCatalog.csv">CSV 数据目录</a></p>
<p>空间图、曲线、分布与对比图：01_figures_hd，PNG ≥1200 DPI。完整数值表：03_data，无需 Origin。</p><ul>{notes}</ul>{cards}''', encoding='utf-8')


def nested(d, path):
    for k in path.split('.'):
        if not isinstance(d, dict): return None
        d = d.get(k)
    return d


SUMMARY_FIELDS = {
    'poses_count': 'poses', 'duration_s': 'duration_s', 'path_length_map_units': 'path_length_map_units',
    'aligned_path_length_m': 'aligned_path_length_m',
    'ate_rmse_m': 'accuracy.ate_m.rmse', 'ate_mean_m': 'accuracy.ate_m.mean', 'ate_p95_m': 'accuracy.ate_m.p95',
    'rpe_trans_rmse_m': 'accuracy.rpe.translation_m.rmse', 'rpe_rot_rmse_deg': 'accuracy.rpe.rotation_deg.rmse',
    'rpe_actual_interval_s': 'accuracy.rpe.actual_interval_s.mean',
    'segment_trans_percent': 'accuracy.distance_segments.translation_percent.mean',
    'segment_rot_deg_per_m': 'accuracy.distance_segments.rotation_deg_per_m.mean',
    'segment_count': 'accuracy.distance_segments.count', 'estimated_coverage_ratio': 'accuracy.estimated_coverage',
    'gt_coverage_ratio': 'accuracy.gt_coverage', 'alignment': 'accuracy.alignment', 'scale': 'accuracy.scale',
    'keyframes_count': 'keyframes', 'map_points_count': 'cloud.count',
    'map_mean_distance_m': 'cloud.reference_error.symmetric_mean_distance_m',
    'map_precision_ratio': 'cloud.reference_error.precision', 'map_recall_ratio': 'cloud.reference_error.recall',
    'map_fscore_ratio': 'cloud.reference_error.fscore', 'map_threshold_m': 'cloud.reference_error.threshold_m',
}
FRAME_FIELDS = {
    'frames_count': 'frames.count', 'valid_pose_ratio': 'frames.valid_pose_ratio',
    'track_mean_ms': 'frames.tracking_ms.mean', 'track_median_ms': 'frames.tracking_ms.median',
    'track_p95_ms': 'frames.tracking_ms.p95', 'track_p99_ms': 'frames.tracking_ms.p99',
    'loss_events_count': 'tracking_losses.count', 'unrecovered_count': 'tracking_losses.unrecovered',
    'overrun_ratio': 'frames.track_exceeds_input_interval_ratio',
}


def comparison(reports, failures, output, formats, dpi, excel=True, charts=False):
    book = Report(output, formats, dpi, excel=excel)
    headers = ['sequence', 'run_id', 'map_id', 'status']+list(SUMMARY_FIELDS)+list(FRAME_FIELDS)
    headers += ['kitti_camera', 'rpe_target_s', 'rpe_tolerance_s', 'association_tolerance_s', 'source', 'sensor_mode']
    modes = sorted({r.metrics.get('sensor_mode', 'unknown') for r in reports} | {f.get('sensor_mode', 'unknown') for f in failures}) or ['unknown']
    entries, grouped = [], {(mode, f'{i:02d}'): [] for mode in modes for i in range(11)}
    for report in reports:
        seq = report.metrics.get('sequence') or 'unknown'
        mode = report.metrics.get('sensor_mode', 'unknown')
        maps = [(k, v) for k, v in report.metrics.items() if k.startswith('map_') and isinstance(v, dict)]
        for stem, metric in maps or [('none', {})]:
            row = [seq, report.path.name, stem, 'NO_ALIGNMENT' if 'accuracy_unavailable' in metric else ('OK' if 'accuracy' in metric else 'NO_GT')]
            row += [nested(metric, key) for key in SUMMARY_FIELDS.values()]
            row += [nested(report.metrics, key) for key in FRAME_FIELDS.values()]
            opts = report.metrics.get('options', {})
            row += [opts.get('kitti_camera'), opts.get('rpe_seconds'), opts.get('rpe_tolerance'), opts.get('max_time_diff'), report.metrics.get('source'), mode]
            entries.append(row)
            if (mode, seq) in grouped: grouped[mode, seq].append(row)
    summary = []
    for (mode, seq), rows in grouped.items():
        if len(rows) == 1: summary.append(rows[0])
        else:
            failed = any(f.get('sequence') == seq and f.get('sensor_mode', 'unknown') == mode for f in failures)
            status = 'MULTIPLE_RUNS_OR_MAPS' if rows else ('FAILED' if failed else 'MISSING')
            row = [seq, None, None, status]+[None]*(len(headers)-4)
            row[-1] = mode
            summary.append(row)
    book.add('SequenceSummary', headers, summary, 'Categorical column / grouped column', 'sequence',
             'ate_rmse_m; rpe_trans_rmse_m; track_mean_ms; valid_pose_ratio',
             '00–10 slots. Multiple runs/maps are in RunSummary; no automatic averaging. Use separate axes for different units.')
    book.add('RunSummary', headers, entries)
    experiment_headers = ['sequence', 'run_id', 'sensor_mode', 'maps_count', 'maps_with_accuracy', 'maps_without_alignment'] + list(FRAME_FIELDS) + ['alignment', 'kitti_camera', 'source']
    book.add('ExperimentSummary', experiment_headers,
             ([r.metrics.get('sequence'), r.path.name, r.metrics.get('sensor_mode', 'unknown'),
               sum(k.startswith('map_') and isinstance(v, dict) for k, v in r.metrics.items()),
               sum(k.startswith('map_') and isinstance(v, dict) and 'accuracy' in v for k, v in r.metrics.items()),
               sum(k.startswith('map_') and isinstance(v, dict) and 'accuracy_unavailable' in v for k, v in r.metrics.items()),
               *[nested(r.metrics, key) for key in FRAME_FIELDS.values()], r.metrics.get('options', {}).get('align'),
               r.metrics.get('options', {}).get('kitti_camera'), r.metrics.get('source')] for r in reports),
             'Categorical column', 'sequence; sensor_mode', 'track_mean_ms; valid_pose_ratio; maps_count',
             'One row per experiment: performance counts once even when it exports multiple independent maps.')
    def base_label(mode, seq):
        return f'{mode}_seq{seq}' if len(modes) > 1 else f'seq{seq}'
    def empty_arrays():
        return {base_label(mode, seq): [] for (mode, seq), rows in grouped.items() if len(rows) <= 1}
    sources = []
    for report in reports:
        seq = report.metrics.get('sequence')
        mode = report.metrics.get('sensor_mode', 'unknown')
        for stem in sorted(k for k in report.metrics if k.startswith('map_')) or ['none']:
            if (mode, seq) in grouped:
                base = base_label(mode, seq)
                label = base if len(grouped[mode, seq]) == 1 else f'{base}_{report.path.name}_{stem}'
                sources.append((label, report, stem))
    # Preserve all samples for box plots; ECDF columns share neither timestamps nor ranks across runs.
    specs = [('ATE', '_absolute_errors', 5, 'm'), ('RPE_trans', '_relative_errors', 2, 'm'),
             ('RPE_rot', '_relative_errors', 3, 'deg'), ('Orientation', '_absolute_errors', 6, 'deg'),
             ('Track', 'tracking_time', 1, 'ms')]
    for title, suffix, col, unit in specs:
        arrays = empty_arrays()
        seen = set()
        for label, report, stem in sources:
            if title == 'Track':
                if report.path.name in seen: continue
                seen.add(report.path.name)
                seq = report.metrics.get('sequence')
                mode = report.metrics.get('sensor_mode', 'unknown')
                label = base_label(mode, seq)
                if sum(r.metrics.get('sequence') == seq and r.metrics.get('sensor_mode', 'unknown') == mode for r in reports) > 1:
                    label += '_'+report.path.name
            key = suffix if title == 'Track' else stem+suffix
            data = report.tables.get(key, {}).get('rows', [])
            arrays[label] = [row[col] for row in data if row[col] is not None]
        arrays = dict(sorted(arrays.items()))
        size = max([len(a) for a in arrays.values()]+[1])
        book.add(title+'_Box', [f'{k}_{unit}' for k in arrays],
                 ([a[i] if i < len(a) else None for a in arrays.values()] for i in range(size)),
                 'Box / violin', 'sequence (column labels)', 'each numeric column', 'No padding with zero; distributions describe samples, not independent trials.')
        ordered = {k: sorted(a) for k, a in arrays.items()}
        book.add(title+'_CDF', [c for k in ordered for c in (f'{k}_{unit}', f'{k}_CDF_ratio')],
                 ([v for a in ordered.values() for v in (a[i] if i < len(a) else None, (i+1)/len(a) if i < len(a) else None)] for i in range(size)),
                 'Multi-curve ECDF', 'each *_'+unit, 'matching *_CDF_ratio', 'Each sequence uses its own X/Y pair; no common time axis.')
    # Shared physical segment length is valid for sequence overlay.
    for column, title, unit in ((4, 'SegmentTranslation', 'percent'), (5, 'SegmentRotation', 'deg_per_m')):
        arrays = empty_arrays()
        for label, report, stem in sources:
            arrays[label] = report.tables.get(stem+'_distance_segments', {}).get('rows', [])
        arrays = dict(sorted(arrays.items()))
        rows = []
        for length in range(100, 801, 100):
            row = [length]
            for a in arrays.values():
                vals = [x[column] for x in a if x[2] == length]
                row += [float(np.mean(vals)) if vals else None, float(np.std(vals)) if vals else None, len(vals) if vals else None]
            rows.append(row)
        book.add(title, ['length_m']+[c for k in arrays for c in (f'{k}_mean_{unit}', f'{k}_std_{unit}', f'{k}_count')], rows,
                 'Multi-curve line / dispersion bars', 'length_m', '*_mean_'+unit,
                 'STD is segment dispersion, not confidence interval; overlapping segments are correlated.')
    book.add('Failures', ['sequence', 'run', 'error', 'sensor_mode'], ([f.get('sequence'), f['run'], f['error'], f.get('sensor_mode', 'unknown')] for f in failures))
    book.notes = ['未完成序列留空；多次实验/多地图不合并，详见 RunSummary。',
                  '不同序列路线不同，不将原始轨迹或绝对时间曲线叠加为精度排名。',
                  '比较误差时检查相同 SE3/Sim3、相机、RPE 间隔、覆盖率；比较耗时还需相同硬件/软件配置。',
                  '单目默认 Sim3 消除未知全局尺度；双目/RGB-D 默认 SE3。sensor_mode 列区分模式，混合批次按模式分别预留 00–10。',
                  'ExperimentSummary 按实验统计性能；RunSummary 按地图列出误差。NO_ALIGNMENT 表示地图过短或退化，误差留空。',
                  '图片由程序直接绘制；完整数值和 X/Y 列见 03_data 与 PlotCatalog。可选 Excel 同样保存全精度数值快照。']
    book.metrics = {'source': 'Per-experiment metrics and full sample tables from this batch',
                    'experiments': [r.metrics for r in reports], 'failures': failures}
    if charts:
        from diagnostic_plots import batch_plots
        batch_plots(book, reports)
    book.finish('SLAM 实验分析 · 00–10 序列对比')
    src = book.path/'02_origin/origin_data.xlsx'
    if excel: src.rename(book.path/'02_origin/sequence_comparison.xlsx')
    html_path = book.path/'report.html'
    html_path.write_text(html_path.read_text().replace('origin_data.xlsx', 'sequence_comparison.xlsx'))
    return book
